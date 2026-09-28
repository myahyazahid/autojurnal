"""Pustaka cek tambahan: tabel/gambar sebagai objek, kebersihan naskah, penulis, dan mutu referensi."""
from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX

from app.engine.docmodel import DocModel
from app.engine.ekstrak import ekstrak_template
from app.engine.objek import garis_tabel
from app.engine.periksa import periksa

PETUNJUK = (
    "Tabel hanya menggunakan garis horizontal. Lebar tabel mengikuti lebar halaman (AutoFit Window). "
    "Gambar diatur in line with text dengan resolusi minimal 300 dpi. Tabel dan gambar wajib diberi keterangan sumber. "
    "Persamaan ditulis dengan Equation Editor. Naskah ditulis dalam bahasa Indonesia. "
    "Naskah tidak boleh menggunakan catatan kaki. Judul tidak boleh mengandung singkatan. "
    "Abstrak tidak boleh mengandung sitasi. Sitasi wajib menggunakan Mendeley atau Zotero. "
    "Setiap referensi wajib mencantumkan DOI. Minimal 80% referensi berasal dari jurnal atau prosiding. "
    "Referensi tidak boleh bersumber dari Wikipedia atau blog."
)


def _png(lebar: int, tinggi: int) -> bytes:
    def potong(jenis: bytes, isi: bytes) -> bytes:
        return struct.pack(">I", len(isi)) + jenis + isi + struct.pack(">I", zlib.crc32(jenis + isi) & 0xFFFFFFFF)

    baris = b"".join(b"\x00" + b"\x80\x80\x80" * lebar for _ in range(tinggi))
    return (b"\x89PNG\r\n\x1a\n" + potong(b"IHDR", struct.pack(">IIBBBBB", lebar, tinggi, 8, 2, 0, 0, 0))
            + potong(b"IDAT", zlib.compress(baris)) + potong(b"IEND", b""))


def _garis(t, **sisi):
    b = OxmlElement("w:tblBorders")
    for nama in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{nama}")
        e.set(qn("w:val"), sisi.get(nama, "none"))
        b.append(e)
    t._tbl.tblPr.append(b)


def _lebar_kolom(t, twip: int):
    for g in t._tbl.tblGrid.findall(qn("w:gridCol")):
        g.set(qn("w:w"), str(twip))


def _cari(doc, awal: str):
    return next(p for p in doc.paragraphs if p.text.startswith(awal))


def _paragraf_sebelum(doc, jangkar, teks: str = ""):
    p = doc.add_paragraph(teks)
    jangkar._p.addprevious(p._p)
    return p


def _gambar_sebelum(doc, jangkar, lebar_px: int, lebar_cm: float, tinggi_cm: float):
    doc.add_picture(io.BytesIO(_png(lebar_px, max(1, lebar_px // 2))), width=Cm(lebar_cm), height=Cm(tinggi_cm))
    p = doc.paragraphs[-1]
    jangkar._p.addprevious(p._p)
    return p


def _run_xml(p, tag: str, **attr):
    r = OxmlElement("w:r")
    e = OxmlElement(tag)
    for k, v in attr.items():
        e.set(qn(f"w:{k}"), v)
    r.append(e)
    p._p.append(r)


def _template_lengkap(template_docx: Path, tmp_path: Path) -> Path:
    doc = docx.Document(str(template_docx))
    t = doc.tables[0]
    _garis(t, top="single", bottom="single", insideH="single")
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _paragraf_sebelum(doc, _cari(doc, "ABSTRAK"), "Email: penulis@contoh.ac.id")
    _paragraf_sebelum(doc, _cari(doc, "KESIMPULAN"), PETUNJUK)
    path = tmp_path / "template_lengkap.docx"
    doc.save(path)
    return path


def test_aturan_tambahan_dibaca_dari_template(template_docx, tmp_path):
    prof, _ = ekstrak_template(str(_template_lengkap(template_docx, tmp_path)))
    tg, ref = prof.tabel_gambar, prof.referensi
    assert (tg.garis_tabel, tg.perataan_tabel, tg.tabel_selebar_halaman) == ("horizontal", "tengah", True)
    assert (tg.gambar_sebaris, tg.min_dpi_gambar, tg.wajib_sumber, tg.persamaan_editor) == (True, 300, True, True)
    assert (prof.naskah.bahasa, prof.naskah.catatan_kaki_dilarang) == ("indonesia", True)
    assert prof.judul.tanpa_singkatan and prof.abstrak.tanpa_sitasi and prof.penulis.wajib_email
    assert (ref.manajer_referensi, ref.wajib_doi, ref.persen_sumber_primer) == ("wajib", True, 80.0)
    assert {"wikipedia", "blog"} <= set(ref.sumber_terlarang)
    assert any("Aturan tambahan yang terbaca" in c for c in prof.catatan_ekstraksi)


def test_template_polos_tidak_menghasilkan_aturan_tanpa_garis(template_docx):
    prof, _ = ekstrak_template(str(template_docx))  # contoh tabelnya tanpa garis sama sekali
    assert prof.tabel_gambar.garis_tabel is None


def test_garis_tabel_dari_style_dan_sel(tmp_path):
    doc = docx.Document()
    grid = doc.add_table(rows=3, cols=3, style="Table Grid")
    horizontal = doc.add_table(rows=3, cols=3, style="Table Grid")
    _garis(horizontal, top="single", bottom="single", insideH="single")  # menimpa garis vertikal dari style
    sel = doc.add_table(rows=3, cols=3)
    for tc in sel._tbl.iter(qn("w:tc")):  # hanya garis atas-bawah per sel
        tcpr = tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for nama in ("top", "bottom"):
            e = OxmlElement(f"w:{nama}")
            e.set(qn("w:val"), "single")
            b.append(e)
        tcpr.append(b)
    path = tmp_path / "garis.docx"
    doc.save(path)
    dm = DocModel(str(path))
    assert [garis_tabel(dm, t).pola for t in dm.tabel] == ["grid", "horizontal", "horizontal"]


def test_naskah_melanggar_aturan_tambahan(template_docx, tmp_path):
    prof, _ = ekstrak_template(str(_template_lengkap(template_docx, tmp_path)))
    doc = docx.Document(str(template_docx))
    doc.paragraphs[0].text = "SISTEM PENDUKUNG KEPUTUSAN (SPK) PEMILIHAN LOKASI USAHA"
    abstrak = _cari(doc, "ABSTRAK")
    abstrak.runs[0].text = "ABSTRAK: Penelitian ini melanjutkan model sebelumnya (Anwar, 2021) dengan data baru. " + abstrak.runs[0].text[9:]

    # tabel 1: grid penuh dan sempit; tabel 2: melewati margin; tabel 3: disisipkan sebagai gambar
    t1 = doc.tables[0]
    t1.style = doc.styles["Table Grid"]
    _lebar_kolom(t1, 800)
    kesimpulan = _cari(doc, "KESIMPULAN")
    _paragraf_sebelum(doc, kesimpulan, "Tabel 2. Data Besar")
    t2 = doc.add_table(rows=3, cols=3)
    _garis(t2, top="single", bottom="single", insideH="single")
    _lebar_kolom(t2, 4000)
    for i, c in enumerate(t2._cells):
        c.text = f"nilai {i}"
    kesimpulan._p.addprevious(t2._tbl)
    _paragraf_sebelum(doc, kesimpulan, "Tabel 3. Tangkapan Layar Data")
    _gambar_sebelum(doc, kesimpulan, 1200, 10, 5)
    _paragraf_sebelum(doc, kesimpulan, "Tabel 2 dan Tabel 3 menunjukkan data  uji yang dipakai. " * 2)

    melayang = _gambar_sebelum(doc, kesimpulan, 40, 12, 6)  # 40 piksel pada 12 cm: resolusi sangat rendah
    inline = melayang._p.find(".//" + qn("wp:inline"))
    inline.tag = qn("wp:anchor")
    inline.find(qn("wp:extent")).addnext(OxmlElement("wp:wrapSquare"))
    _paragraf_sebelum(doc, kesimpulan, "Gambar 1. Grafik Hasil")
    persamaan = _gambar_sebelum(doc, kesimpulan, 300, 5, 0.8)
    persamaan.add_run("\t(2)")

    kotor = _paragraf_sebelum(doc, kesimpulan, "Hasil ini diuji ulang. ")
    biru = kotor.add_run("Bagian ini berwarna biru. ")
    biru.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
    kuning = kotor.add_run("Bagian ini masih disorot. ")
    kuning.font.highlight_color = WD_COLOR_INDEX.YELLOW
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "91")
    ins.set(qn("w:author"), "Penulis")
    r = OxmlElement("w:r")
    teks = OxmlElement("w:t")
    teks.text = "Kalimat sisipan."
    r.append(teks)
    ins.append(r)
    kotor._p.append(ins)
    _run_xml(kotor, "w:commentReference", id="5")
    _run_xml(kotor, "w:footnoteReference", id="1")
    _paragraf_sebelum(doc, kesimpulan)
    _paragraf_sebelum(doc, kesimpulan)

    pustaka = doc.paragraphs[-1]
    for ref in ("Wikipedia. (2020). Sistem informasi. Diakses dari https://id.wikipedia.org/wiki/Sistem_informasi",
                "Budi, C. (2019). Buku sistem informasi. Penerbit Andi."):
        p = doc.add_paragraph(ref)
        pustaka._p.addnext(p._p)
    path = tmp_path / "naskah_lanjut.docx"
    doc.save(path)

    temuan, _ = periksa(DocModel(str(path)), prof)
    kode = {t.kode for t in temuan}
    harap = {
        "tabel_gambar.garis", "tabel_gambar.lebar_tabel", "tabel_gambar.melebihi_margin", "tabel_gambar.tabel_berupa_gambar",
        "tabel_gambar.tanpa_sumber", "tabel_gambar.gambar_melayang", "tabel_gambar.resolusi_rendah",
        "tabel_gambar.persamaan_gambar", "tabel_gambar.nomor_persamaan",
        "naskah.track_changes", "naskah.komentar_lama", "naskah.sorotan", "naskah.teks_berwarna", "naskah.spasi_ganda",
        "naskah.baris_kosong", "naskah.catatan_kaki", "judul.singkatan", "abstrak.sitasi", "penulis.email",
        "referensi.manajer", "referensi.tanpa_doi", "referensi.sumber_primer", "referensi.sumber_terlarang",
    }
    assert harap <= kode, harap - kode
    assert "naskah.bahasa" not in kode
    pesan = {t.kode: t.pesan for t in temuan}
    assert pesan["tabel_gambar.garis"].startswith("Tabel 1 memakai grid penuh")
    assert "SPK" in pesan["judul.singkatan"] and "(Anwar, 2021)" in pesan["abstrak.sitasi"]
    assert "Wikipedia" in pesan["referensi.sumber_terlarang"] and "kuning" in pesan["naskah.sorotan"]
    assert "biru (#2F5496)" in pesan["naskah.teks_berwarna"]
    assert next(t for t in temuan if t.kode == "referensi.manajer").tingkat == "wajib"
    assert "$" not in " ".join(t.pesan for t in temuan)


def test_jumlah_halaman_tidak_percaya_metadata_basi(tmp_path, monkeypatch):
    from app.engine import docmodel
    from app.engine.halaman import jumlah_halaman

    doc = docx.Document()
    for i in range(3):
        p = doc.add_paragraph(f"Paragraf halaman {i + 1}. " * 5)
        if i:
            p.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    path = tmp_path / "halaman.docx"
    doc.save(path)
    monkeypatch.setattr(docmodel.DocModel, "_baca_halaman", staticmethod(lambda sumber: 1))  # metadata basi
    assert jumlah_halaman(DocModel(str(path))) == (3, "render")

    polos = docx.Document()
    for _ in range(120):
        polos.add_paragraph("Kalimat panjang untuk mengisi halaman naskah uji tanpa penanda render dari Word. " * 4)
    path2 = tmp_path / "polos.docx"
    polos.save(path2)
    n, dasar = jumlah_halaman(DocModel(str(path2)))
    assert dasar == "perkiraan" and n > 5
