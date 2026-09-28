"""Pengecil ukuran berkas: buang font sematan, kompres gambar, dan endpoint kecilkan/unduh."""
from __future__ import annotations

import io
import os
import zipfile

import docx
from docx.shared import Cm
from lxml import etree

from app.engine.kecilkan import kecilkan, rincian

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _sematkan_font(data: bytes, ukuran: int) -> bytes:
    """Tiru opsi Word "Embed fonts in the file": part font + rujukannya di fontTable dan settings."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in zin.infolist():
            isi = zin.read(i.filename)
            if i.filename == "word/fontTable.xml":
                root = etree.fromstring(isi)
                f = root.find(f"{{{W}}}font")
                e = etree.SubElement(f, f"{{{W}}}embedRegular")
                e.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rIdF1")
                isi = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif i.filename == "[Content_Types].xml":
                isi = isi.replace(b"</Types>", b'<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/></Types>')
            elif i.filename == "word/settings.xml":
                root = etree.fromstring(isi)
                root.insert(0, etree.Element(f"{{{W}}}embedTrueTypeFonts"))
                isi = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            z.writestr(i, isi)
        z.writestr("word/fonts/font1.odttf", os.urandom(ukuran))  # acak = tidak bisa dikompres
        z.writestr("word/_rels/fontTable.xml.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rIdF1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font1.odttf"/></Relationships>')
    return out.getvalue()


def _naskah_besar(tmp_path, font_kb: int = 2500) -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Naskah uji dengan font sematan. " * 20)
    import test_pustaka

    doc.add_picture(io.BytesIO(test_pustaka._png(2400, 1200)), width=Cm(12))
    b = io.BytesIO()
    doc.save(b)
    return _sematkan_font(b.getvalue(), font_kb * 1024)


def test_kecilkan_buang_font_sematan(tmp_path):
    data = _naskah_besar(tmp_path)
    awal = rincian(data)
    assert awal["total_kb"] > 1900 and awal["font_kb"] >= 2400
    baru, lap = kecilkan(data, 1900)
    assert lap["tercapai"] and lap["akhir_kb"] < 1900 and len(lap["langkah"]) == 1  # cukup tahap tanpa kehilangan kualitas
    z = zipfile.ZipFile(io.BytesIO(baru))
    assert not [n for n in z.namelist() if n.startswith("word/fonts/")]
    assert b"embedTrueTypeFonts" not in z.read("word/settings.xml") and b"embedRegular" not in z.read("word/fontTable.xml")
    d = docx.Document(io.BytesIO(baru))
    assert len(d.inline_shapes) == 1 and d.paragraphs[0].text.startswith("Naskah uji")


def test_kecilkan_gambar_bila_masih_besar(tmp_path):
    data = _naskah_besar(tmp_path, font_kb=10)
    baru, lap = kecilkan(data, target_kb=5)  # target sangat kecil memaksa tahap kompres gambar
    assert len(lap["langkah"]) > 1 and lap["akhir_kb"] < rincian(data)["total_kb"]
    d = docx.Document(io.BytesIO(baru))
    assert len(d.inline_shapes) == 1


def test_api_kecilkan_dan_unduh(tmp_path):
    from fastapi.testclient import TestClient

    from app.main import app
    from test_komentar import _admin

    naskah = tmp_path / "besar.docx"
    naskah.write_bytes(_naskah_besar(tmp_path))
    with TestClient(app) as c:
        _admin(c)
        j = c.get("/api/jurnal").json()[0]
        with open(naskah, "rb") as f:
            r = c.post("/api/cek", files={"naskah": ("besar.docx", f)}, data={"jurnal_id": j["id"]}).json()
        assert r["ukuran"]["total_kb"] > r["ukuran"]["maks_kb"] == 1900 and r["file_kecil_tersedia"] is False
        k = c.post(f"/api/cek/{r['id']}/kecilkan").json()
        assert k["kecil"]["tercapai"] and k["kecil"]["akhir_kb"] < 1900 and k["file_kecil_tersedia"]
        unduh = c.get(f"/api/cek/{r['id']}/unduh?kecil=true")
        assert unduh.status_code == 200 and "_DICEK_kecil.docx" in unduh.headers["content-disposition"]
        assert len(unduh.content) < 1900 * 1024
        assert list(docx.Document(io.BytesIO(unduh.content)).comments)  # komentar hasil cek tetap ada
