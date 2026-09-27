"""Dokumen sintetis untuk pengujian — tidak memakai naskah asli."""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

import pytest

# data uji terpisah — harus di-set sebelum app.config diimpor
os.environ.setdefault("AUTOJURNAL_DATA", tempfile.mkdtemp(prefix="autojurnal-tes-"))
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def _p(doc, teks, font="Times New Roman", pt=12, bold=False, italic=False, align=None, spasi=1.0, indent=None):
    p = doc.add_paragraph()
    r = p.add_run(teks)
    r.font.name, r.font.size, r.bold, r.italic = font, Pt(pt), bold, italic
    if align:
        p.alignment = align
    p.paragraph_format.line_spacing = spasi
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    return p


def _tabel(doc, font="Times New Roman", pt=10):
    t = doc.add_table(rows=3, cols=3)
    for i, row in enumerate(t.rows):
        for j, c in enumerate(row.cells):
            r = c.paragraphs[0].add_run(f"sel {i}{j}")
            r.font.name, r.font.size = font, Pt(pt)
    return t


ISI = ("Penelitian ini membahas penerapan metode yang relevan dengan masalah nyata di lapangan. "
       "Data dikumpulkan dari berbagai sumber yang dapat dipertanggungjawabkan. ")


@pytest.fixture
def template_docx(tmp_path) -> Path:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.right_margin = Cm(3)
    s.left_margin = Cm(4)
    C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
    _p(doc, "JUDUL ARTIKEL (font Times New Roman, 14 pt, Bold, maksimal 15 kata)", pt=14, bold=True, align=C)
    _p(doc, "Nama Penulis (Times New Roman 10 pt)", pt=10, align=C)
    _p(doc, "Afiliasi Penulis", pt=10, align=C)
    _p(doc, "ABSTRAK: Abstrak berisi tujuan, metode, dan hasil penelitian. Abstrak terdiri atas 150 s.d. 250 kata. (font Times New Roman, 11 pt)", pt=11, align=J)
    _p(doc, "Kata Kunci: terdiri atas 3 s.d. 5 kata kunci, dipisahkan dengan titik koma (;). (Times New Roman 11 pt)", pt=11)
    for judul, isi in [
        ("PENDAHULUAN", "Pendahuluan berisi latar belakang dan kesenjangan penelitian. " + ISI),
        ("METODE PENELITIAN", "Metode penelitian harus menjelaskan desain penelitian dan teknik analisis data. " + ISI),
        ("HASIL DAN PEMBAHASAN", "Judul tabel diletakkan di atas tabel. Judul gambar diletakkan di bawah gambar. " + ISI),
    ]:
        _p(doc, judul, bold=True)
        _p(doc, isi, align=J, indent=1.0)
        _p(doc, "*) Teks utama ditulis dengan font Times New Roman 12 pt dan spasi 1. Setiap paragraf tidak boleh terdiri atas satu kalimat.", align=J, indent=1.0)
    _p(doc, "Tabel 1. Contoh Tabel", pt=10, align=C)
    _tabel(doc)
    _p(doc, "KESIMPULAN", bold=True)
    _p(doc, "Kesimpulan menjawab tujuan penelitian. " + ISI, align=J, indent=1.0)
    _p(doc, "DAFTAR PUSTAKA", bold=True)
    _p(doc, "Daftar pustaka minimal 10 referensi, 80% terbit dalam 10 tahun terakhir, ditulis dengan gaya APA dan diurutkan sesuai abjad.", align=J)
    _p(doc, "Anwar, B. (2021). Contoh referensi. Jurnal Contoh, 1(2), 1-10.", align=J)
    path = tmp_path / "template.docx"
    doc.save(path)
    return path


@pytest.fixture
def naskah_docx(tmp_path) -> Path:
    """Naskah dengan pelanggaran yang disengaja."""
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)  # Letter, bukan A4
    s.top_margin = s.bottom_margin = s.right_margin = s.left_margin = Cm(2.54)
    C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
    _p(doc, "ANALISIS PENGARUH MEDIA SOSIAL TERHADAP MINAT BELAJAR MAHASISWA TINGKAT AKHIR DI PERGURUAN TINGGI SWASTA KOTA SURABAYA TAHUN 2025",
       pt=14, bold=True, align=C)
    _p(doc, "Budi Santoso", pt=10, align=C)
    _p(doc, "Universitas Contoh", pt=10, align=C)
    _p(doc, "ABSTRAK: " + "Penelitian ini singkat sekali. " * 10, pt=11, align=J)
    _p(doc, "Kata Kunci: media sosial, minat belajar, mahasiswa, survei, regresi, kuantitatif, Surabaya", pt=11)
    _p(doc, "PENDAHULUAN", bold=True)
    _p(doc, "Media sosial berkembang pesat (Anwar, 2021). " + ISI, font="Calibri", align=J, indent=1.0)
    _p(doc, "Paragraf ini hanya satu kalimat dan tidak punya kalimat penjelas sama sekali di dalamnya.", align=J, indent=1.0)
    _p(doc, "HASIL DAN PEMBAHASAN", bold=True)
    _p(doc, "Hasil penelitian ditunjukkan sebagai berikut menurut Citra (2019). " + ISI, align=J, indent=1.0, spasi=1.5)
    _tabel(doc)
    _p(doc, "Tabel 1. Hasil Uji", pt=10, align=C)
    _p(doc, "KESIMPULAN", bold=True)
    _p(doc, "Kesimpulan penelitian ini jelas. " + ISI, align=J, indent=1.0)
    _p(doc, "DAFTAR PUSTAKA", bold=True)
    for ref in [
        "Zulkifli, A. (2005). Buku lama sekali. Penerbit Lama.",
        "Anwar, B. (2021). Contoh referensi. Jurnal Contoh, 1(2), 1-10.",
        "Dewi, C. (2010). Referensi yang tidak disitasi. Jurnal Lain, 3(1), 5-9.",
    ]:
        _p(doc, ref, align=J)
    path = tmp_path / "naskah.docx"
    doc.save(path)
    return path
